"""Generator arkusza z metoda babilonska.

Tworzy metoda-babilonska.xlsx - te sama metode co w Pythonie i C++, tyle ze
w arkuszu kalkulacyjnym, gdzie kazda iteracja jest osobnym wierszem i widac
ja golym okiem.

UWAGA O NAZWACH FUNKCJI. Excel przechowuje formuly wewnetrznie po angielsku
i dopiero wyswietla je w jezyku interfejsu. openpyxl zapisuje to, co dostanie,
wiec w kodzie musza byc nazwy angielskie (SQRT, ABS, LOG10, IFERROR).
W polskim Excelu uzytkownik zobaczy PIERWIASTEK, MODUL.LICZBY, LOG10
i JEZELI.BLAD - i to jest poprawne zachowanie, nie blad.

Uruchomienie:

    python generuj_arkusz.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WYJSCIE = Path(__file__).resolve().parent / "metoda-babilonska.xlsx"

LICZBA_ITERACJI = 10

AKCENT = "9C90FC"
TLO_NAGLOWKA = "EEEAE4"


def obramowanie_dolne():
    return Border(bottom=Side(style="thin", color="D8D3CB"))


def buduj():
    skoroszyt = Workbook()
    arkusz = skoroszyt.active
    arkusz.title = "Metoda babilonska"

    arkusz["A1"] = "Metoda babilonska (metoda Herona)"
    arkusz["A1"].font = Font(bold=True, size=14)
    arkusz["A2"] = "x(k+1) = 0,5 * ( x(k) + n / x(k) )"
    arkusz["A2"].font = Font(italic=True, color="6D7173")

    # Dane wejsciowe - jedyne dwie komorki, ktore uzytkownik zmienia.
    arkusz["A4"] = "n (liczba podpierwiastkowa)"
    arkusz["B4"] = 2
    arkusz["A5"] = "x0 (punkt startowy)"
    arkusz["B5"] = 1
    arkusz["A6"] = "wzorzec: PIERWIASTEK(n)"
    arkusz["B6"] = "=SQRT($B$4)"

    for komorka in ("A4", "A5", "A6"):
        arkusz[komorka].font = Font(bold=True)
    for komorka in ("B4", "B5"):
        arkusz[komorka].fill = PatternFill("solid", fgColor=AKCENT)
        arkusz[komorka].font = Font(bold=True, color="FFFFFF")
    arkusz["B6"].number_format = "0.000000000000000"

    naglowki = ["k", "x(k)", "n / x(k)", "blad bezwzgledny", "poprawne cyfry"]
    wiersz_naglowka = 8
    for kolumna, tekst in enumerate(naglowki, start=1):
        komorka = arkusz.cell(row=wiersz_naglowka, column=kolumna, value=tekst)
        komorka.font = Font(bold=True, size=9)
        komorka.fill = PatternFill("solid", fgColor=TLO_NAGLOWKA)
        komorka.alignment = Alignment(horizontal="center")
        komorka.border = obramowanie_dolne()

    pierwszy = wiersz_naglowka + 1
    for k in range(LICZBA_ITERACJI + 1):
        w = pierwszy + k
        arkusz.cell(row=w, column=1, value=k)

        if k == 0:
            arkusz.cell(row=w, column=2, value="=$B$5")
        else:
            # Serce metody: srednia poprzedniego przyblizenia i tego,
            # co ono dzieli w liczbie n.
            arkusz.cell(row=w, column=2, value=f"=0.5*(B{w - 1}+$B$4/B{w - 1})")

        arkusz.cell(row=w, column=3, value=f"=IFERROR($B$4/B{w},\"\")")
        arkusz.cell(row=w, column=4, value=f"=ABS(B{w}-$B$6)")
        # Liczba poprawnych cyfr znaczacych = -log10(blad wzgledny).
        # Gdy blad spadnie do zera, logarytm nie istnieje - wtedy "dokladny".
        arkusz.cell(
            row=w, column=5,
            value=f'=IFERROR(-LOG10(D{w}/$B$6),"dokladny")',
        )

        for kolumna in range(2, 5):
            arkusz.cell(row=w, column=kolumna).number_format = "0.000000000000000"
        arkusz.cell(row=w, column=5).number_format = "0.00"
        for kolumna in range(1, 6):
            arkusz.cell(row=w, column=kolumna).border = obramowanie_dolne()

    arkusz.column_dimensions["A"].width = 28
    for kolumna in range(2, 6):
        arkusz.column_dimensions[get_column_letter(kolumna)].width = 22

    ostatni = pierwszy + LICZBA_ITERACJI
    arkusz.cell(
        row=ostatni + 2, column=1,
        value="Zmien B4 i B5 - cala tabela przeliczy sie sama.",
    ).font = Font(italic=True, color="6D7173")
    arkusz.cell(
        row=ostatni + 3, column=1,
        value="Kolumna 'poprawne cyfry' ma sie w kazdym kroku mniej wiecej "
              "podwajac - to jest zbieznosc kwadratowa.",
    ).font = Font(italic=True, color="6D7173")

    arkusz.freeze_panes = "A9"
    skoroszyt.save(WYJSCIE)
    return WYJSCIE


if __name__ == "__main__":
    sciezka = buduj()
    print(f"Zapisano: {sciezka}")
